"""Выгрузка в Excel (три листа) и CSV."""
from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Change, Product

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF")
RED_FILL = PatternFill("solid", fgColor="FDE2E2")
GREEN_FILL = PatternFill("solid", fgColor="DCFCE7")


def _sheet(wb, title: str, headers: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def to_xlsx(products: list[Product], changes: list[Change], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cur = products[0].currency if products else ""
    money = f'#,##0.00 "{cur}"' if cur else "#,##0.00"
    wb = Workbook()
    wb.remove(wb.active)

    # 1) все товары
    ws = _sheet(wb, "Товары", ["Название", "Категория", "Цена", "В наличии, шт.", "Рейтинг", "Артикул (UPC)", "Ссылка"],
                [52, 20, 12, 15, 11, 19, 14])
    for p in products:
        ws.append([p.title, p.category, p.price, p.stock, "★" * p.rating + "☆" * (5 - p.rating), p.upc, "открыть"])
        link = ws.cell(ws.max_row, 7)
        link.hyperlink, link.style = p.url, "Hyperlink"
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = money
    if products:
        ws.auto_filter.ref = f"A1:G{ws.max_row}"
        ws.conditional_formatting.add(f"D2:D{ws.max_row}", CellIsRule(operator="equal", formula=["0"], fill=RED_FILL))

    # 2) сводка по категориям
    groups: dict[str, list[Product]] = defaultdict(list)
    for p in products:
        groups[p.category or "Без категории"].append(p)
    ws2 = _sheet(wb, "По категориям", ["Категория", "Товаров", "Средняя цена", "Мин. цена", "Макс. цена",
                                       "Всего на складе", "Средний рейтинг"], [24, 10, 14, 12, 12, 16, 16])
    for cat, items in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        prices = [p.price for p in items]
        ws2.append([cat, len(items), round(sum(prices) / len(prices), 2), min(prices), max(prices),
                    sum(p.stock for p in items), round(sum(p.rating for p in items) / len(items), 1)])
    for row in ws2.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = money

    # 3) изменения с прошлого запуска
    ws3 = _sheet(wb, "Изменения", ["Что случилось", "Название", "Было", "Стало", "Разница", "Ссылка"],
                 [22, 52, 12, 12, 12, 14])
    if not changes:
        ws3.append(["Изменений нет — или это первый запуск, сравнивать не с чем"])
    for c in changes:
        delta = (c.new - c.old) if c.kind == "price" else None
        ws3.append([c.label, c.title, c.old, c.new, delta, "открыть"])
        link = ws3.cell(ws3.max_row, 6)
        link.hyperlink, link.style = c.url, "Hyperlink"
        fill = RED_FILL if c.kind in ("stock_out", "removed") or (delta or 0) > 0 else GREEN_FILL
        ws3.cell(ws3.max_row, 1).fill = fill
        if c.kind in ("price", "new", "removed"):        # здесь в «было/стало» цены
            for col in (3, 4, 5):
                ws3.cell(ws3.max_row, col).number_format = money
        else:                                            # а здесь — остаток в штуках
            for col in (3, 4):
                ws3.cell(ws3.max_row, col).number_format = '0" шт."'

    wb.save(path)
    return path


def to_csv(products: list[Product], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = ["title", "category", "price", "currency", "stock", "rating", "upc", "url"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:     # utf-8-sig — чтобы Excel открыл кириллицу
        w = csv.DictWriter(f, fieldnames=fields, delimiter=";", extrasaction="ignore")
        w.writeheader()
        for p in products:
            w.writerow(p.as_dict())
    return path
