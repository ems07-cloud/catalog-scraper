import csv

from openpyxl import load_workbook

from scraper.export import to_csv, to_xlsx
from scraper.history import History, diff
from scraper.models import Product


def prod(upc, price=10.0, stock=5, title=None, category="Poetry", rating=3):
    return Product(url=f"https://x/{upc}", title=title or f"Книга {upc}", category=category, price=price,
                   currency="£", stock=stock, rating=rating, upc=upc)


def test_first_run_has_nothing_to_compare():
    assert diff({}, [prod("a")]) == []


def test_all_kinds_of_changes():
    before = [prod("a", 10), prod("b", 20), prod("c", 30, stock=4), prod("d", 40, stock=0), prod("e", 50)]
    now = [prod("a", 10), prod("b", 18.5), prod("c", 30, stock=0), prod("d", 40, stock=7), prod("f", 60)]
    prev = {p.upc: {"title": p.title, "url": p.url, "price": p.price, "stock": p.stock} for p in before}
    got = {(c.kind, c.upc, c.old, c.new) for c in diff(prev, now)}
    assert got == {("price", "b", 20, 18.5), ("stock_out", "c", 4, 0), ("stock_back", "d", 0, 7),
                   ("new", "f", None, 60), ("removed", "e", 50, None)}


def test_history_keeps_last_run_per_source(tmp_path):
    h = History(tmp_path / "h.sqlite")
    h.save("site-1", [prod("a", 10)])
    h.save("site-1", [prod("a", 12)])
    h.save("site-2", [prod("z", 99)])
    assert h.last_snapshot("site-1")["a"]["price"] == 12
    assert list(h.last_snapshot("site-2")) == ["z"]
    assert h.last_snapshot("unknown") == {}


def test_excel_has_three_sheets_with_data(tmp_path):
    products = [prod("a", 10, category="Poetry", rating=5), prod("b", 30, stock=0, category="Travel")]
    changes = diff({"a": {"title": "Книга a", "url": "https://x/a", "price": 12, "stock": 5}}, products)
    wb = load_workbook(to_xlsx(products, changes, tmp_path / "out.xlsx"))
    assert wb.sheetnames == ["Товары", "По категориям", "Изменения"]

    goods = list(wb["Товары"].iter_rows(values_only=True))
    assert goods[1][:5] == ("Книга a", "Poetry", 10, 5, "★★★★★")
    assert wb["Товары"]["G2"].hyperlink.target == "https://x/a"

    cats = {r[0]: r for r in wb["По категориям"].iter_rows(min_row=2, values_only=True)}
    assert cats["Travel"][1:6] == (1, 30, 30, 30, 0)

    changes_rows = list(wb["Изменения"].iter_rows(min_row=2, values_only=True))
    assert ("Цена снизилась", "Книга a", 12, 10, -2) == changes_rows[0][:5]


def test_csv_opens_in_excel_with_cyrillic(tmp_path):
    path = to_csv([prod("a", title="Книга с кириллицей")], tmp_path / "out.csv")
    assert path.read_bytes()[:3] == b"\xef\xbb\xbf"           # BOM для Excel
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig"), delimiter=";"))
    assert rows[0]["title"] == "Книга с кириллицей" and rows[0]["price"] == "10.0"


def test_changes_sheet_formats_money_and_pieces(tmp_path):
    prev = {"a": {"title": "a", "url": "u", "price": 12, "stock": 5}, "b": {"title": "b", "url": "u", "price": 9, "stock": 0}}
    products = [prod("a", 10), prod("b", 9, stock=3), prod("c", 7)]
    ws = load_workbook(to_xlsx(products, diff(prev, products), tmp_path / "o.xlsx"))["Изменения"]
    formats = {ws.cell(r, 1).value: ws.cell(r, 4).number_format for r in range(2, ws.max_row + 1)}
    assert "£" in formats["Цена снизилась"] and "£" in formats["Новый товар"]
    assert formats["Снова в наличии"] == '0" шт."'
